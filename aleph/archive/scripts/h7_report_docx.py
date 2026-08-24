"""Build the H.7 full-cell single-cell-integration report (.docx).

Consolidates the H.7 reboot session (Phase A/B, GATE-B, the two lamellipodium
geometries, the overnight production rounds, research/novelty) into one
presentable Word report + the PI decision list. Mirrors layer2_report_docx.py.

Usage:  python -m aleph.scripts.h7_report_docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

_FIGS = Path(__file__).resolve().parents[1] / "outputs" / "h7" / "figs"
_OUT = Path(__file__).resolve().parents[1] / "outputs" / "h7" / "H7_Report.docx"


def _h(doc, text, level):
    doc.add_heading(text, level=level)


def _fig(doc, name, caption):
    path = _FIGS / name
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        run = cap.add_run(caption)
        run.italic = True
        run.font.size = Pt(9)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(f"[figure {name} not found]")


def _bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def build() -> Path:
    doc = Document()

    title = doc.add_heading("H.7 — Full Single-Cell Integration (MCF7): GATE-B Report", level=0)
    sub = doc.add_paragraph()
    r = sub.add_run("Fine-grained mechanistic HOOMD assembly of a complete physiological MCF7 cell; "
                    "emergent cortical tension (KU-3.5) + single-cell spreading.  2026-06-07  ·  "
                    "branch h7/full-cell-integration  ·  PROVISIONAL (no KU-3.5 conclusion drawn).")
    r.italic = True
    r.font.size = Pt(10)

    # ---- Executive summary ----
    _h(doc, "1. Executive summary", 1)
    doc.add_paragraph(
        "The full physiological MCF7 cell now assembles and runs at the full ×40 mesoscopic "
        "scale (≈1000 effective filaments) in one HOOMD run — cortex + crosslinkers + bipolar "
        "myosin + nucleus + osmotic turgor + cytoplasm viscosity + membrane + focal-adhesion "
        "clutch, all at physiological setpoints. Cortical tension γ is measured as three "
        "SEPARATE channels (active method-of-planes, rigid M-SHAKE Lagrange, passive turgor "
        "Young-Laplace), with the turgor channel never folded into a total."
    )
    p = doc.add_paragraph()
    p.add_run("Headline finding (honest, no gate-loosening): ").bold = True
    p.add_run(
        "the emergent cortical tension floors ~6–10× BELOW the [0.35, 0.65] mN/m reference band, "
        "and this holds across seeds, both integrators, turnover on/off, and the full ×40 scale "
        "range. No clean emergent channel reaches the band. The root cause is a TIMESCALE GAP: "
        "the tension-generating processes (myosin contraction, actin turnover τ½≈10 s) are "
        "seconds-scale, while feasible fine-grained MD reaches only ~milliseconds — so the "
        "active-tension steady state is unreachable. This reproduces and sharpens the project's "
        "γ-floor result inside the complete physiological cell."
    )

    # ---- The cell ----
    _h(doc, "2. The assembled cell (full ×40 production scale)", 1)
    _fig(doc, "h7_baseline_cell_components.png",
         "Fig 1. The physiological-baseline MCF7 cell at full ×40 scale (1000 effective filaments): "
         "actin cortex shell, crosslinkers, bipolar myosin minifilaments, myosin heads, nucleus; "
         "plasma-membrane shell + turgor + cytoplasm as field compartments. Built via the single "
         "manifest (configs/mcf7_baseline.yaml) → the unified Cell.build.")
    doc.add_paragraph(
        "Physiological setpoints (all literature-anchored, no magic numbers): R_cell 7.5 µm "
        "(Wagner 2011); cytoplasm η 65.9 Pa·s (Hu 2024); turgor Π₀ 133 Pa (band-implied, Stewart "
        "2011 lineage); nucleus E 4.7 kPa (KU-3.B2.1); membrane γ 0.10 mN/m (KU-3.B1)."
    )

    # ---- GATE-B findings ----
    _h(doc, "3. GATE-B — emergent cortical tension (the 3 channels)", 1)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text = "channel"; hdr[1].text = "clean value"; hdr[2].text = "reading"
    rows = [
        ("γ_active (myosin, MOP)", "0.0001 ± 0 mN/m (n=8)",
         "floors at ~0; turnover-ON does not raise it; scale-independent"),
        ("γ_rigid (free cortex, M-SHAKE)", "0.06–0.09 mN/m (n=3, ×40-converged)",
         "turgor implies 0.50 but transmits only ~⅛ to the backbone — below band"),
        ("γ_passive (turgor Young-Laplace)", "0.499 mN/m",
         "CIRCULAR identity with the turgor setpoint — not an independent measurement"),
        ("reference band", "[0.35, 0.65] mN/m",
         "rounded/de-adhered non-MCF7 proxy (see Decision 3)"),
    ]
    for a, b, c in rows:
        cells = tbl.add_row().cells
        cells[0].text = a; cells[1].text = b; cells[2].text = c
    doc.add_paragraph()
    doc.add_paragraph(
        "Why the active channel floors — a timescale gap (confirmed): turnover-ON left γ_active at "
        "0.0001 because τ½=10 s ≫ the ms run; and γ_active≈0 in BOTH the constrained and "
        "unconstrained integrators → this is a real generation/timescale limitation, not an "
        "M-SHAKE artifact. The lone ever-in-band number (0.391, FA-adhered) was shown to be an "
        "integrin-overload artifact (the dynamic integrin–ligand bond was wired r0=0), and is "
        "discarded. The ×40 coarse-graining convergence study confirms both the active floor and "
        "the structural value are mesh-converged, not coarse-graining artifacts."
    )

    # ---- Spreading ----
    _h(doc, "4. Single-cell spreading — two lamellipodium geometries", 1)
    doc.add_paragraph(
        "Both single-cell leading-edge geometries were implemented and build in the full cell "
        "(membrane-load OFF, step-1; the Bieling/Funk dendritic mechanism unchanged): "
        "basal_ring (isotropic spreading rim) and polarized_patch (migrating leading edge). They "
        "are compared by basal footprint A/A0 so the right one can be chosen against your "
        "experimental A/A0 data."
    )
    _fig(doc, "h7_spreading_compare_full.png",
         "Fig 2. Spreading comparison (full ×40): basal_ring (isotropic, footprint +24%) vs "
         "polarized_patch (directional, polarization 0.90). A cell-movement animation over time is "
         "at outputs/h7/figs/h7_spreading_movement.gif.")

    # ---- Novelty / method ----
    _h(doc, "5. Novelty + method positioning", 1)
    doc.add_paragraph(
        "An independent novelty review (citations verified) places the contribution at the "
        "INTEGRATION level (narrow but real): a whole single cell built explicit-everything "
        "(cortex + bipolar myosin + catch-bond crosslinkers + integrin clutch + nucleus + turgor "
        "+ cytoplasm + membrane) in one dynamical run, plus the 3-channel γ measurement discipline "
        "(turgor never folded into a total). The decisive risk it flags — that the emergent γ is "
        "“a promise, not a result” — is exactly what this session quantified: at accessible "
        "fine-grained timescales the active γ does not reach the band. The conventional active-gel "
        "/ active-shell continuum sidesteps this by working at the seconds-scale directly."
    )

    # ---- Decisions ----
    _h(doc, "6. DECISIONS NEEDED (PI)", 1)
    p = doc.add_paragraph()
    p.add_run("None of these were forced autonomously. ").bold = True
    p.add_run("Decision 1 is the pivotal scientific call.")
    _bullets(doc, [
        "1. Active-γ path forward — (a) ACCELERATED DYNAMICS (raise myosin/turnover rates to reach "
        "steady state in feasible steps) — but this breaks Hill-validity and would produce a "
        "Hill-invalid in-band number that could be misread as success, so it needs an explicit "
        "validity gate (I deliberately did NOT run it overnight); (b) MULTISCALE active-gel γ-seam "
        "(hand the coarse-grained stress to a seconds-scale continuum); (c) RE-TARGET — a spread "
        "adherent MCF7 is traction/stress-fibre dominated, so the right observable may be traction, "
        "not cortical γ.",
        "2. Lamellipodium geometry — basal_ring (isotropic) vs polarized_patch (migrating). Both "
        "built; choose by matching your experimental A/A0.",
        "3. Reference band — [0.35,0.65] mN/m is a rounded/de-adhered non-MCF7 proxy; confirm the "
        "correct MCF7 target (this is a gate-contract change).",
        "4. γ_passive — anchor turgor to an independent osmotic datum (to de-circularize the passive "
        "channel) or keep labeling it a setpoint.",
        "5. fa.py force-free-integrin fix — implement it (delicate multi-file core-FA change; spec'd "
        "+ test-gated) to get a clean FA-adhered γ_rigid? Not blocking (clean γ_rigid already via "
        "the free-cortex path).",
        "6. Citation fix — validation/pereverzev.py KU-2.18 has the wrong title for Bangasser 2013 "
        "(metadata drift); recommend a SourceEvidence title fix.",
    ])

    # ---- Deliverables ----
    _h(doc, "7. Deliverables (all committed on h7/full-cell-integration)", 1)
    _bullets(doc, [
        "Single physiological-baseline manifest (configs/mcf7_baseline.yaml) + loader → the unified "
        "Cell.build (Phase A: assembler unification + manifest).",
        "GATE-B pipeline (scripts/h7_gate_b.py) + the 3-channel γ estimator (cortex/cortical_tension.py).",
        "6 GPU production rounds on the RTX A5000 (γ ensembles, turnover test, clean γ_rigid, ×40 "
        "convergence); curated summaries in outputs/h7/production/H7_*.md.",
        "Both lamellipodium geometries + A/A0 comparison + cell-movement animation "
        "(outputs/h7/figs/h7_spreading_movement.gif).",
        "3 research reviews (direction-review, novelty-analysis, 16 verified paper candidates) in "
        "docs/v2_audit/ and references/.",
        "Stability map + the fa.py force-free-integrin fix specification.",
    ])

    _OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(_OUT))
    return _OUT


if __name__ == "__main__":
    print(f"[h7-report] wrote {build()}")
